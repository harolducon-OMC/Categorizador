# -*- coding: utf-8 -*-
"""
Prueba de humo SIN llamar a la API: usa un LLM simulado para validar toda la lógica.
Ejecutar desde la raíz del proyecto:  python tests/smoke_test.py
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd

from core import brands, coding, export, insights
from core.awareness import build_analysis, category_table, BANKS, BANK_ALIASES, BANK_NORMALIZATIONS


class FakeLLM:
    fast, smart = "fake-fast", "fake-smart"

    def json(self, system, user, model=None, max_tokens=2000):
        if "construir un libro de códigos" in system:
            return {"categorias": [
                {"nombre": "Durabilidad", "descripcion": "Dura mucho"},
                {"nombre": "Precio", "descripcion": "Costo / valor"},
                {"nombre": "Otros", "descripcion": "x"},
            ]}
        if "Asigna cada respuesta" in system:
            items = json.loads(user.split("Respuestas a clasificar:\n")[1])
            out = []
            for it in items:
                t, cats = it["texto"].lower(), []
                if "dura" in t:
                    cats.append("durabilidad")        # minúsculas a propósito: debe mapear
                if "preci" in t or "barat" in t:
                    cats.append("Precio")
                out.append({"id": it["id"], "categorias": cats or ["Categoría inventada"]})
            return {"resultados": out}
        if "MARCAS" in system:
            items = json.loads(user.split("Respuestas:\n")[1])
            found = {"bancolombia": "Bancolombia", "bancolonbia": "Bancolombia", "davivienda": "Davivienda", "bbva": "BBVA"}
            return {"resultados": [{"id": it["id"], "marcas": [v for k, v in found.items() if k in it["texto"].lower()]
                                    or ([] if "ninguna" in it["texto"].lower() else [])} for it in items]}
        if "insights" in system:
            data = json.loads(user.split("Datos por categoría:\n")[1])
            return {"resumen": "Lectura global.", "categorias": [
                {"categoria": d["categoria"], "explicacion": "Dicen X porque Y.", "citas": [d["respuestas"][0]["id"], 99999]}
                for d in data]}
        raise AssertionError("prompt no reconocido")


llm = FakeLLM()
texts = (["Me gusta porque dura mucho tiempo"] * 3 + ["Es barato, buen precio"] * 2
         + ["Dura y además tiene buen precio", "no sé", "", "Me atiende bien el personal"])

# --- categorías
frame = coding.discover_categories(llm, "¿Por qué?", texts)
assert [c["nombre"] for c in frame] == ["Durabilidad", "Precio"], frame      # "Otros" se descarta
labels, stats = coding.classify(llm, "¿Por qué?", texts, frame, workers=2)
assert labels[0] == ["Durabilidad"] and labels[5] == ["Durabilidad", "Precio"]
assert labels[6] == [coding.NSNR_LABEL] and labels[7] == [coding.NSNR_LABEL]
assert labels[8] == [coding.OTHER]                                          # categoría inventada -> Otros
assert stats["unique_texts"] == 4 and stats["unclassified"] == 0, stats     # deduplicación
table = coding.frequency_table(labels, [c["nombre"] for c in frame])
assert coding.valid_base(labels) == 7
assert abs(table.set_index("Categoría").loc["Durabilidad", "%"] - 4 / 7) < 1e-9
seg = pd.Series(["H", "M", "H", "M", "H", "M", "H", "M", "H"])
matrix, bases = coding.crosstab_by_segment(labels, seg, ["Durabilidad", "Precio"])
assert set(matrix.columns) == {"H", "M"} and sum(bases.values()) == 7
print("OK categorías:", table.to_dict("records"))

# --- marcas
bt = ["Bancolombia", "bancolonbia", "Davivienda y BBVA", "ninguna", "no sé", "Davivienda"]
raw = brands.extract_mentions(llm, "¿Qué banco recuerda?", bt, workers=2)
mapping = brands.build_mapping(raw, 85)
mp = dict(zip(mapping["Mención"], mapping["Marca normalizada"]))
assert mp["Bancolombia"] == "Bancolombia"
ents, aliases = brands.entities_and_aliases(raw, mp)
assert set(ents) == {"Bancolombia", "Davivienda", "BBVA"} and "bancolombia" in aliases["Bancolombia"], (ents, aliases)
print("OK marcas:", ents)

# --- insights + excel
ins = insights.build_insights(llm, "¿Por qué?", texts, labels, ["Durabilidad", "Precio", coding.OTHER])
assert all(len(c["citas"]) == 1 for c in ins["categorias"])                 # cita 99999 inventada se descarta
df = pd.DataFrame({"id": range(len(texts)), "Pregunta": texts})
xlsx = export.build_report_excel(df, {"¿Por qué?": {"labels": labels, "categories": ["Durabilidad", "Precio", coding.OTHER]}},
                                 {"¿Por qué?": ins})
assert xlsx[:2] == b"PK"
print("OK insights/excel:", len(xlsx), "bytes")

# --- awareness (lógica original)
d = pd.DataFrame({"TOM": ["Bancolombia", "bbva", ""], "ESP": ["Davivienda", "", "Bancolombia"], "sexo": ["H", "M", "H"]})
cfg = {"name": "t", "tom": "TOM", "esp": "ESP", "ayu": "0"}
try:
    _out, ind, _raw = build_analysis(d, {}, cfg, BANKS, BANK_ALIASES, BANK_NORMALIZATIONS, 0)
    t = category_table(ind, d, None, None, BANKS).set_index("Entidad")
    assert t.loc["Bancolombia", "TOM"] == 1 and t.loc["Bancolombia", "Espontaneo"] == 1 and t.loc["BBVA", "TOM"] == 1
    assert t.loc["Davivienda", "Espontaneo"] == 1
    print("OK awareness:", t.loc[["Bancolombia", "Davivienda", "BBVA"], ["TOM", "Espontaneo", "AWA %"]].to_dict("index"))
except KeyError as exc:  # noqa: BLE001
    raise
print("TODO OK")

# --- validación: segunda opinión + codificación manual
import io
import random
from core import validation

random.seed(3)
phr = {
    "Durabilidad": ["dura mucho", "muy resistente", "no se daña", "dura años", "aguanta bastante"],
    "Precio": ["es barato", "buen precio", "precio justo", "económico", "cuesta poco"],
    "Comodidad": ["muy cómodo", "cómoda para caminar", "suave y liviana", "se siente cómodo"],
}
pre = ["me gusta porque", "la verdad", "creo que", "para mí", "en general", ""]
post = ["", "sinceramente", "la verdad", "y ya", "como siempre", "gracias"]
validation_texts, validation_labels = [], []
for _ in range(300):
    selected = random.sample(list(phr), random.choice([1, 1, 2]))
    validation_texts.append(
        f"{random.choice(pre)} " + " y ".join(random.choice(phr[c]) for c in selected)
        + f" {random.choice(post)}"
    )
    validation_labels.append(selected)
true_labels = [list(labels) for labels in validation_labels]
wrong = random.sample(range(300), 12)
for i in wrong:
    validation_labels[i] = ["Precio"] if validation_labels[i] != ["Precio"] else ["Durabilidad"]

validation_result = validation.second_opinion(validation_texts, validation_labels, list(phr))
metrics = validation_result["metrics"].set_index("Categoría")
assert (metrics["Kappa"] > 0.6).all(), metrics
caught = set(validation_result["review"]["Fila"]) & set(wrong)
assert len(caught) >= 8, (len(caught), validation_result["flagged"])
assert validation_result["flagged"] < 60, validation_result["flagged"]
print("OK segunda opinión: kappa", metrics["Kappa"].to_dict(), "| errores detectados", len(caught), "de 12")

manual_xlsx = validation.sample_for_manual(validation_texts, list(phr), n=80)
manual_sheet = pd.read_excel(io.BytesIO(manual_xlsx), sheet_name="Codificar")
assert list(manual_sheet.columns) == ["Fila", "Respuesta", "Codigo manual"] and len(manual_sheet) == 80
manual_sheet["Codigo manual"] = [" | ".join(true_labels[int(row)]).lower() for row in manual_sheet["Fila"]]
manual, unknown = validation.parse_manual(manual_sheet, list(phr))
assert len(manual) == 80 and not unknown
manual_table, manual_total = validation.compare_with_manual(validation_labels, manual, list(phr))
assert manual_total["n"] == 80 and 0.8 < manual_total["exact"] < 1.0, manual_total
print("OK codificación manual: exactitud", round(manual_total["exact"], 2))

validation_xlsx = export.build_report_excel(
    pd.DataFrame({"Respuesta": validation_texts}),
    {"Pregunta": {"labels": validation_labels, "categories": list(phr)}},
    {},
    {"Pregunta": {"second": validation_result}},
)
from openpyxl import load_workbook
assert "Validación" in load_workbook(io.BytesIO(validation_xlsx), read_only=True).sheetnames
print("OK reporte: hoja Validación incluida")

# --- BERTopic: integración comprobada con dependencias simuladas (sin descargas)
import sys
import types
from unittest.mock import patch
from core import topics

class FakeSentenceTransformer:
    def __init__(self, model_name):
        self.model_name = model_name
    def encode(self, docs, **kwargs):
        return [[1.0, 0.0] for _ in docs]

class FakeBERTopic:
    def __init__(self, **kwargs):
        self.docs = []
    def fit_transform(self, docs, embeddings):
        self.docs = list(docs)
        return [0] * len(docs), None
    def get_topic_info(self):
        return pd.DataFrame([{"Topic": 0, "Count": len(self.docs)}])
    def get_representative_docs(self, topic_id):
        return self.docs[:5]
    def get_topic(self, topic_id):
        return [("calidad", 1.0), ("precio", 0.7)]

fake_bertopic = types.ModuleType("bertopic")
fake_bertopic.BERTopic = FakeBERTopic
fake_sentence_transformers = types.ModuleType("sentence_transformers")
fake_sentence_transformers.SentenceTransformer = FakeSentenceTransformer

class FakeTopicNamer:
    smart = "mock-model"
    def json(self, system, user, model=None, max_tokens=3000):
        return {"categorias": [{"nombre": "Calidad del producto", "descripcion": "Opiniones sobre calidad."}]}

with patch.dict(sys.modules, {
    "bertopic": fake_bertopic,
    "sentence_transformers": fake_sentence_transformers,
}):
    topic_frame = topics.discover_with_bertopic(
        FakeTopicNamer(), "¿Qué opinas?", [f"respuesta distinta número {i}" for i in range(60)], 5
    )
assert topic_frame[0]["nombre"] == "Calidad del producto", topic_frame
print("OK BERTopic: agrupamiento y nombrado de categorías integrados")
